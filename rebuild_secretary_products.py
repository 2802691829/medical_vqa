from __future__ import annotations

import re
import shutil
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import jieba
import jieba.posseg as pseg
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path.home() / "xwechat_files" / "wxid_6p586amrjhon22_4737" / "msg" / "file" / "2026-07"
YEAR = "2025"
SUBJECT = "秘书学"
CATEGORY = "学术期刊"
JOURNAL = "秘书之友"

STOPWORDS = {
    "的", "了", "在", "和", "与", "及", "或", "为", "是", "有", "也", "都", "而", "其", "中", "对",
    "以", "从", "由", "把", "被", "将", "并", "等", "这", "那", "一个", "一种", "通过", "进行",
    "不断", "同时", "以及", "对于", "方面", "工作", "相关", "可以", "需要", "必须", "我们", "他们",
    "自己", "这个", "这些", "这种", "由于", "因此", "并且", "如果", "作为", "其中", "主要", "进一步",
    "实现", "推进", "提升", "加强", "具有", "形成", "推动", "做好", "问题", "方式", "情况", "过程",
    "内容", "作用", "要求", "发展", "建设", "研究", "秘书之友",
}

DOMAIN_TERMS = [
    "秘书学", "秘书工作", "办公室", "三服务", "公文写作", "公文处理", "督查督办", "调查研究",
    "会议服务", "文稿写作", "信息工作", "档案管理", "党政秘书", "教学秘书", "科研秘书",
    "行政管理", "基层治理", "人工智能", "生成式人工智能", "DeepSeek", "AIGC", "智慧办公",
    "电子公文", "请示报告", "办文指数", "党政机关", "高校办公室", "高职院校", "应用文写作",
    "党政公文", "文电工作", "文秘写作", "秘书实践", "秘书素养", "文稿服务", "公务接待",
    "会议报到", "会议摄影", "视频会议", "讲话稿", "发言稿", "咨政报告", "申论归纳概括题",
    "基层减负", "第一议题", "三步工作法", "四轮驱动", "四定四化四融合", "清单制",
    "标准操作规范", "BOPPPS模型", "金课标准", "课程思政", "政产学研用", "研究性教学",
    "办公自动化", "流动办公", "知识管理", "数字化转型", "AI赋能", "AI幻觉",
    "人机协同", "认知升维", "提示词", "大语言模型", "电子公文办理",
    "党委信息", "督查工作", "督办机制", "督办工作", "复核审查", "结果运用",
]

ARTICLE_SEPARATORS = ["-", "——", "──"]


def find_secretary_dir() -> Path:
    dirs = [p for p in ROOT.iterdir() if p.is_dir()]
    return max(dirs, key=lambda p: len(list(p.glob("*.txt"))))


def parse_title_author(path: Path) -> tuple[str, str]:
    stem = path.stem
    for sep in ARTICLE_SEPARATORS:
        if sep in stem:
            left, right = stem.rsplit(sep, 1)
            if right.strip():
                return left.strip(), right.strip()
    return stem.strip(), ""


def author_from_content(raw: str) -> str:
    for line in raw.splitlines()[:20]:
        stripped = normalize_fragment(line)
        if stripped.startswith("作者："):
            value = stripped.replace("作者：", "", 1).strip()
            if value:
                return value
        if stripped.startswith("□"):
            value = stripped.replace("□", "", 1).strip()
            value = re.sub(r"\s+", "", value)
            if 1 < len(value) <= 20:
                return "、".join(re.findall(r"[\u4e00-\u9fff]{2,4}", value)) or value
    return ""


def source_tail_number(raw: str) -> str:
    lines = [line.strip() for line in raw.splitlines() if line.strip()]
    tail = " ".join(lines[-10:])
    nums = re.findall(r"\d{1,4}", tail)
    return nums[-1] if nums else ""


def normalize_fragment(text: str) -> str:
    text = text.replace("\ufeff", "").replace("\u3000", " ")
    text = re.sub(r"[ \t]+", " ", text.strip())
    text = re.sub(r"[■□◆◇●○�]", "", text)
    text = re.sub(r"(?<=[\u4e00-\u9fff])\s*[-－—]\s*(?=[\u4e00-\u9fff])", "", text)
    text = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", text)
    text = re.sub(r"(?<=\d)\s+(?=\d)", "", text)
    text = re.sub(r"\s+([，。；：！？、）》】])", r"\1", text)
    text = re.sub(r"([（【《])\s+", r"\1", text)
    text = text.replace("Chat GPT", "ChatGPT")
    text = text.replace("ArtificialIntelligence", "Artificial Intelligence")
    text = text.replace("AI 写作", "AI写作")
    text = text.replace("AI 工具", "AI工具")
    text = text.replace("AI 技术", "AI技术")
    text = text.replace("AI 大模型", "AI大模型")
    return text.strip()


def is_header_footer(text: str) -> bool:
    compact = re.sub(r"\s+", "", text)
    if not compact:
        return True
    if compact in {"学习与修养", "秘书工作", "公文写作", "调查研究", "会务工作", "档案工作", "督查工作"}:
        return True
    if re.search(r"秘书之友\s*2025[·.]\d+", text):
        return True
    if re.search(r"^\d+\s*秘书之友\s*2025[·.]\d+$", text):
        return True
    return False


def split_columns(line: str) -> list[str]:
    raw = line.rstrip()
    if not raw.strip() or is_header_footer(raw):
        return []
    gaps = [(m.start(), m.end()) for m in re.finditer(r"[ \t]{4,}", raw)]

    def choose(target: int, after: int = 0):
        candidates = [
            gap for gap in gaps
            if gap[0] >= after and target - 10 <= gap[1] <= target + 10 and gap[1] - gap[0] >= 4
        ]
        if not candidates:
            return None
        return min(candidates, key=lambda gap: abs(gap[1] - target))

    g1 = choose(32)
    g2 = choose(62, g1[1] if g1 else 0)
    if g1 and g2:
        parts = [raw[:g1[0]], raw[g1[1]:g2[0]], raw[g2[1]:]]
    elif g1:
        parts = [raw[:g1[0]], raw[g1[1]:]]
    else:
        parts = re.split(r"[ \t]{14,}", raw)

    return [normalize_fragment(part) for part in parts if normalize_fragment(part)]


def should_start_body(line: str) -> bool:
    compact = re.sub(r"\s+", "", line)
    if len(compact) < 15:
        return False
    return bool(re.search(r"[，。；：！？]|^[一二三四五六七八九十]+、|引言|结语", compact))


def clean_generic(raw: str, title: str, author: str) -> str:
    if "秘书之友" not in raw and "参考文献" not in raw and title in raw[:200]:
        body = raw
    else:
        pages: list[list[list[str]]] = []
        cols = [[], [], []]
        ref_skip = [False, False, False]
        started = False
        skip_table = False
        skip_black_block = False

        for line in raw.splitlines():
            stripped = line.strip()
            if not stripped:
                skip_table = False
                continue
            if re.search(r"[■□◆◇]{8,}", stripped):
                skip_black_block = True
                continue
            if skip_black_block:
                if "秘书之友" in stripped and re.search(r"2025[·.]\d+", stripped):
                    skip_black_block = False
                continue
            if re.match(r"^\s*(表|图)\s*\d+", stripped):
                skip_table = True
                continue
            if skip_table:
                continue
            if is_header_footer(stripped):
                if any(cols):
                    pages.append(cols)
                    cols = [[], [], []]
                    ref_skip = [False, False, False]
                continue
            if not started:
                if should_start_body(stripped):
                    started = True
                else:
                    continue
            parts = split_columns(line)
            for idx, part in enumerate(parts[:3]):
                if idx >= len(ref_skip) or ref_skip[idx]:
                    continue
                compact = re.sub(r"\s+", "", part)
                if "参考文献" in compact or "作者单位" in compact or re.match(r"^［?\d+］?", compact):
                    ref_skip[idx] = True
                    continue
                cols[idx].append(part)
        if any(cols):
            pages.append(cols)

        chunks = []
        for page_cols in pages:
            for col in page_cols:
                text = "".join(col)
                if text:
                    chunks.append(text)
        body = "".join(chunks)

    body = normalize_fragment(body)
    if len(body) < 120:
        fallback_lines = []
        start = False
        for line in raw.splitlines():
            stripped = normalize_fragment(line)
            if not stripped or is_header_footer(stripped):
                continue
            compact = re.sub(r"\s+", "", stripped)
            if compact in {re.sub(r"\s+", "", title), re.sub(r"\s+", "", author), "卷首语"}:
                continue
            if compact.startswith("作者："):
                continue
            if "□" in line and len(compact) <= 20:
                continue
            if not start and should_start_body(stripped):
                start = True
            if start:
                fallback_lines.append(stripped)
        body = normalize_fragment("".join(fallback_lines))

    body = re.sub(rf"^论文名[:：]?{re.escape(title)}", "", body).strip()
    body = re.sub(rf"^{re.escape(title)}", "", body).strip()
    if author:
        body = re.sub(rf"^作者[:：]?{re.escape(author)}", "", body).strip()
        body = re.sub(rf"^{re.escape(author)}", "", body).strip()
    body = re.sub(r"参考文献[:：]?.*$", "", body, flags=re.S)
    body = re.sub(r"（作者单位[:：]?.*?）", "", body)
    body = re.sub(r"作者单位[:：]?.*$", "", body, flags=re.S)
    body = re.sub(r"［\d+］|\[\d+\]", "", body)
    body = re.sub(r"(?m)^\s*(表|图)\s*\d+.*$", "", body)
    body = body.replace("秘书之友", "")
    body = re.sub(r"2025[·.]\d+", "", body)
    body = re.sub(r"\s+", " ", body).strip()
    body = re.sub(r"([。！？])", r"\1\n", body)
    body = re.sub(r"\n{2,}", "\n", body).strip()

    header = [title]
    if author:
        header.append(author)
    else:
        header.append("")
    return "\n".join(header) + "\n\n" + body + "\n"


def tokenize(text: str) -> list[tuple[str, str]]:
    for term in DOMAIN_TERMS:
        jieba.add_word(term, freq=200000, tag="n")
    tokens = []
    for word, flag in pseg.cut(text):
        word = word.strip()
        if len(word) < 2 or word in STOPWORDS:
            continue
        if re.fullmatch(r"[\W_]+|\d+|[A-Za-z]+", word):
            continue
        tokens.append((word, flag))
    return tokens


def write_xlsx(path: Path, tokens: list[tuple[str, str]], source_name: str, tail_num: str) -> None:
    counts = Counter(word for word, _ in tokens)
    pos_counts: dict[str, Counter] = defaultdict(Counter)
    for word, flag in tokens:
        pos_counts[word][flag] += 1

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "词频词性表"
    ws.append(["排名", "词语", "词性标注", "词频", "原文末尾数字", "来源文件"])
    for rank, (word, count) in enumerate(counts.most_common(), 1):
        pos = pos_counts[word].most_common(1)[0][0]
        ws.append([rank, word, pos, count, tail_num, source_name])

    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")
        row[0].alignment = Alignment(horizontal="center")
        row[2].alignment = Alignment(horizontal="center")
        row[3].alignment = Alignment(horizontal="center")
        row[4].alignment = Alignment(horizontal="center")

    widths = [8, 26, 12, 10, 14, 56]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    secretary_dir = find_secretary_dir()
    files = sorted(secretary_dir.glob("*.txt"))
    out_root = secretary_dir / "成品"
    if out_root.exists():
        backup = secretary_dir / f"成品_旧版备份_{datetime.now():%Y%m%d_%H%M%S}"
        shutil.move(str(out_root), str(backup))
        print(f"BACKUP_OUT={backup}")
    out_root.mkdir(parents=True, exist_ok=True)

    mapping_rows = []
    for idx, source in enumerate(files, 1):
        raw = source.read_text(encoding="utf-8", errors="replace")
        title, author = parse_title_author(source)
        author = author_from_content(raw) or author
        tail_num = source_tail_number(raw)
        seq = f"{idx:03d}"
        stem = f"{YEAR}_{SUBJECT}_{CATEGORY}_{JOURNAL}_{seq}"
        out_dir = out_root / stem
        cleaned = clean_generic(raw, title, author)
        txt_path = out_dir / f"{stem}.txt"
        xlsx_path = out_dir / f"{stem}_词频词性表.xlsx"
        out_dir.mkdir(parents=True, exist_ok=True)
        txt_path.write_text(cleaned, encoding="utf-8", newline="\n")
        write_xlsx(xlsx_path, tokenize(cleaned), source.name, tail_num)
        mapping_rows.append((seq, tail_num, source.name, title, author, len(cleaned)))

    print(f"ROOT={secretary_dir}")
    print(f"OUT={out_root}")
    print(f"COUNT={len(files)}")
    for row in mapping_rows[:8]:
        print("MAP", row)
    print("TAIL_DUPLICATES_ARE_PAGE_NUMBERS=YES")


if __name__ == "__main__":
    main()
