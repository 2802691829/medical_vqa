from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(r"C:\Users\28026\Desktop\毕业设计\outputs")
OUT.mkdir(exist_ok=True)


BLUE = "1F4E79"
LIGHT = "D9EAF7"
PALE = "EEF5FB"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def setup(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(1.7)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(1.65)
    sec.right_margin = Cm(1.65)
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for h in ("Heading 1", "Heading 2"):
        styles[h].font.name = "Microsoft YaHei"
        styles[h]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        styles[h].font.color.rgb = RGBColor.from_string(BLUE)


def title(doc, text, sub):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(20)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(sub)
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(90, 90, 90)


def heading(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string(BLUE)


def answer_grid(doc, title_text, answers, per_row=5):
    heading(doc, title_text)
    rows = (len(answers) + per_row - 1) // per_row + 1
    table = doc.add_table(rows=rows, cols=per_row * 2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i in range(per_row):
        set_cell_text(table.cell(0, i * 2), "题号", True, "FFFFFF")
        set_cell_text(table.cell(0, i * 2 + 1), "答案", True, "FFFFFF")
        set_cell_shading(table.cell(0, i * 2), BLUE)
        set_cell_shading(table.cell(0, i * 2 + 1), BLUE)
    for idx, ans in enumerate(answers, 1):
        row = (idx - 1) // per_row + 1
        col = ((idx - 1) % per_row) * 2
        set_cell_text(table.cell(row, col), idx)
        set_cell_text(table.cell(row, col + 1), ans)
        if row % 2 == 1:
            set_cell_shading(table.cell(row, col), PALE)
            set_cell_shading(table.cell(row, col + 1), PALE)
    doc.add_paragraph()


def rows_table(doc, title_text, headers, rows, widths=None):
    heading(doc, title_text)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        set_cell_text(table.cell(0, j), h, True, "FFFFFF")
        set_cell_shading(table.cell(0, j), BLUE)
        if widths:
            table.cell(0, j).width = widths[j]
    for i, row in enumerate(rows, 1):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            set_cell_text(cells[j], value)
            cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j else WD_ALIGN_PARAGRAPH.CENTER
            if i % 2 == 1:
                set_cell_shading(cells[j], PALE)
    doc.add_paragraph()
    return table


def entry_rows(entries):
    rows = []
    for no, lines in entries:
        rows.append((no, "\n".join(lines)))
    return rows


def build_hw7():
    doc = Document()
    setup(doc)
    title(doc, "汪政扬第七次作业", "第五章判断选择题、业务题二、业务题三")
    answer_grid(doc, "一、第五章判断题", ["错", "错", "对", "错", "对", "错", "对", "错", "对", "错"])
    answer_grid(doc, "二、第五章单项选择题", ["A", "B", "C", "C", "C", "A", "B", "A", "B", "D"])
    answer_grid(doc, "三、第五章多项选择题", ["BCE", "BE", "BCD", "BCE", "ADE", "ACE", "BDE"], per_row=4)

    entries2 = [
        ("1", ["借：在途物资 5,150", "借：应交税费-应交增值税(进项税额) 800", "贷：银行存款 5,950"]),
        ("2", ["借：原材料 5,150", "贷：在途物资 5,150"]),
        ("3", ["借：生产成本-A产品 25,000", "借：生产成本-B产品 12,000", "借：制造费用 800", "借：管理费用 300", "贷：原材料 38,100"]),
        ("4", ["借：制造费用 2,000", "借：应交税费-应交增值税(进项税额) 320", "贷：银行存款 2,320"]),
        ("5", ["借：制造费用 560", "贷：其他应收款 500", "贷：库存现金 60"]),
        ("6", ["借：制造费用 600", "贷：库存现金 600"]),
        ("7", ["借：应付职工薪酬 54,000", "贷：银行存款 54,000"]),
        ("8", ["借：生产成本-A产品 10,000", "借：生产成本-B产品 30,000", "借：制造费用 6,000", "借：管理费用 8,000", "贷：应付职工薪酬 54,000"]),
        ("9", ["借：生产成本-A产品 4,150", "借：生产成本-B产品 12,450", "借：制造费用 2,490", "借：管理费用 3,320", "贷：应付职工薪酬 22,410"]),
        ("10", ["借：制造费用 1,250", "贷：累计折旧 1,250"]),
        ("11", ["借：银行存款 58,000", "贷：主营业务收入-A产品 30,000", "贷：主营业务收入-B产品 20,000", "贷：应交税费-应交增值税(销项税额) 8,000", "借：主营业务成本 41,000", "贷：库存商品-A产品 24,000", "贷：库存商品-B产品 17,000"]),
        ("12", ["借：应收账款 3,480", "贷：其他业务收入 3,000", "贷：应交税费-应交增值税(销项税额) 480", "借：其他业务成本 2,800", "贷：原材料 2,800"]),
        ("13", ["借：制造费用 720", "借：应交税费-应交增值税(进项税额) 108", "贷：库存现金 828"]),
        ("14", ["借：预付账款 1,500", "借：应交税费-应交增值税(进项税额) 150", "贷：银行存款 1,650", "借：制造费用 500", "贷：预付账款 500"]),
        ("15", ["制造费用合计 14,920；按生产工人工资比例分配：A产品 3,730，B产品 11,190。", "借：生产成本-A产品 3,730", "借：生产成本-B产品 11,190", "贷：制造费用 14,920"]),
        ("16", ["A产品完工成本 = 25,000 + 10,000 + 4,150 + 3,730 = 42,880", "借：库存商品-A产品 42,880", "贷：生产成本-A产品 42,880"]),
    ]
    rows_table(doc, "四、第五章业务题二：会计分录", ["题号", "答案"], entry_rows(entries2))

    entries3 = [
        ("1", ["借：固定资产 301,200", "借：应交税费-应交增值税(进项税额) 48,120", "贷：银行存款 349,320"]),
        ("2", ["月折旧额 = 301,200 x (1 - 10%) / 10 / 12 = 2,259", "借：制造费用 2,259", "贷：累计折旧 2,259"]),
        ("3", ["账面价值 = 301,200 - 2,259 x 4 = 292,164；减值 = 292,164 - 282,000 = 10,164", "借：资产减值损失 10,164", "贷：固定资产减值准备 10,164"]),
        ("4", ["决定出售但不划归为持有待售资产，不编制会计分录。"]),
        ("5", ["借：固定资产清理 900", "贷：库存现金 900"]),
        ("6", ["借：固定资产清理 282,000", "借：累计折旧 9,036", "借：固定资产减值准备 10,164", "贷：固定资产 301,200", "借：银行存款 322,480", "贷：固定资产清理 278,000", "贷：应交税费-应交增值税(销项税额) 44,480"]),
        ("7", ["净损失 = 282,000 + 900 - 278,000 = 4,900", "借：营业外支出 4,900", "贷：固定资产清理 4,900"]),
    ]
    rows_table(doc, "五、第五章业务题三：固定资产分录", ["题号", "答案"], entry_rows(entries3))
    doc.save(OUT / "汪政扬_第七次作业.docx")


def build_hw8():
    doc = Document()
    setup(doc)
    title(doc, "汪政扬第八次作业", "第七至第十二章客观题、改错账、报表题与业务分录")

    answer_grid(doc, "一、第七章单项选择题", list("DCDDABCCAD"))
    answer_grid(doc, "二、第七章多项选择题", ["ABCD", "ABCD", "ABCD", "BCD", "BC", "CD", "ABCD", "ABC", "ABC", "ABC"])
    answer_grid(doc, "三、第七章判断题", ["对", "对", "对", "错", "对", "对", "错", "错", "错", "错"])

    answer_grid(doc, "四、第八章单项选择题", ["D", "B", "C", "C", "A", "C", "B", "B", "A"])
    answer_grid(doc, "五、第八章多项选择题", ["ABCD", "ABCD", "BC", "A", "AB", "ABD", "ABC", "BC", "ABD", "ABD"])
    answer_grid(doc, "六、第八章判断题", ["对", "对", "错", "错", "对", "错", "错", "错", "对", "对"])

    answer_grid(doc, "七、第九章单项选择题", list("BBCDCCD CAD".replace(" ", "")))
    answer_grid(doc, "八、第九章多项选择题", ["ABCD", "AC", "ACD", "ACD", "BD", "AC", "BCD", "ABCD", "ABC", "AB"])
    answer_grid(doc, "九、第九章判断题", ["错", "错", "错", "错", "错", "对", "错", "对", "错", "对"])

    answer_grid(doc, "十、第十章单项选择题", ["B", "C", "D", "B", "A", "A", "B", "C"])
    answer_grid(doc, "十一、第十章多项选择题", ["AC", "ABD", "AB", "ACD", "ABCD", "ABC"], per_row=3)
    answer_grid(doc, "十二、第十章判断题", ["错", "错", "错", "对", "对", "对", "错", "错"])

    answer_grid(doc, "十三、第十一章单项选择题", ["C", "B", "B", "A", "B", "A", "C", "B", "D", "A", "C", "C", "A", "A"])
    answer_grid(doc, "十四、第十一章多项选择题", ["ABD", "ABC", "AC", "BD", "AC", "ABC", "ABCD", "BCD", "BCD", "ABD", "CD", "BCD", "ABD", "BCD", "ACD"])
    answer_grid(doc, "十五、第十一章判断题", ["对", "对", "对", "对", "错", "对", "错", "错", "错", "错", "对", "错", "对", "对"])

    answer_grid(doc, "十六、第十二章单项选择题", ["A", "B", "A", "B", "C", "D", "C", "A", "D", "A"], per_row=5)
    answer_grid(doc, "十七、第十二章多项选择题", ["AD", "ABD", "AD", "ABCD", "CD", "ACD", "ABCD", "ABCD", "ABCD", "AD"])
    answer_grid(doc, "十八、第十二章判断题", ["错", "错", "对", "错", "对", "错", "错"])

    correction = [
        ("1", "原贷记库存现金错误。红字冲销原分录；再作：借：管理费用 800；贷：银行存款 800。"),
        ("2", "借方科目错误。红字冲销原分录；再作：借：待摊费用 1,800；贷：银行存款 1,800。"),
        ("3", "金额多记 846,000。红字更正：借：库存商品 846,000；贷：生产成本 846,000。"),
        ("4", "科目正确、金额正确，不作更正。"),
        ("5", "过账金额少记 900。采用补充登记法：借：应付账款 900；贷：银行存款 900。"),
    ]
    rows_table(doc, "十九、改错账", ["题号", "答案"], correction)

    rows_table(doc, "二十、第九章业务题二：银行存款余额调节表", ["项目", "金额"], [
        ("企业银行存款日记账余额", "24,000"),
        ("加：银行已收、企业未收", "4,900"),
        ("减：银行已付、企业未付", "3,500"),
        ("调节后余额", "25,400"),
        ("银行对账单余额", "24,400"),
        ("加：企业已收、银行未收", "2,500"),
        ("减：企业已付、银行未付", "1,500"),
        ("调节后余额", "25,400"),
    ])

    rows_table(doc, "二十一、第十章报表题一：利润表", ["项目", "本期金额"], [
        ("营业收入", "3,620,000"),
        ("营业成本", "2,180,000"),
        ("营业税金及附加", "68,000"),
        ("销售费用", "180,000"),
        ("管理费用", "210,000"),
        ("财务费用", "98,000"),
        ("资产减值损失", "169,000"),
        ("投资收益", "167,000"),
        ("营业利润", "882,000"),
        ("营业外收入", "13,000"),
        ("营业外支出", "69,600"),
        ("利润总额", "825,400"),
        ("所得税费用", "267,800"),
        ("净利润", "557,600"),
    ])

    rows_table(doc, "二十二、第十章报表题二：资产负债表主要项目", ["项目", "金额", "项目", "金额"], [
        ("货币资金", "1,896,900", "短期借款", "100,000"),
        ("交易性金融资产", "25,000", "应付票据", "190,000"),
        ("应收票据", "150,000", "应付账款", "1,600,000"),
        ("应收账款", "1,352,500", "应付职工薪酬", "168,000"),
        ("预付账款", "200,000", "应交税费", "367,000"),
        ("其他应收款", "20,000", "应付利息", "8,000"),
        ("存货", "4,996,600", "应付股利", "96,000"),
        ("流动资产合计", "8,641,000", "其他应付款", "30,000"),
        ("长期股权投资", "600,000", "一年内到期的非流动负债", "600,000"),
        ("固定资产", "4,182,000", "流动负债合计", "3,159,000"),
        ("在建工程", "1,126,000", "长期借款", "1,760,000"),
        ("工程物资", "300,000", "负债合计", "4,919,000"),
        ("无形资产", "1,130,000", "股本", "10,000,000"),
        ("非流动资产合计", "7,338,000", "资本公积", "300,000"),
        ("资产总计", "15,979,000", "盈余公积", "280,000"),
        ("", "", "未分配利润", "480,000"),
        ("", "", "所有者权益合计", "11,060,000"),
        ("", "", "负债和所有者权益总计", "15,979,000"),
    ])

    rows_table(doc, "二十三、第十一章业务题一：银行存款余额调节表", ["项目", "金额"], [
        ("企业银行存款日记账余额", "535,000"),
        ("加：银行已收、企业未收", "17,008"),
        ("减：银行已付、企业未付", "3,468"),
        ("调节后余额", "548,540"),
        ("银行对账单余额", "544,885"),
        ("加：企业已收、银行未收", "4,700"),
        ("减：企业已付、银行未付", "1,045"),
        ("调节后余额", "548,540"),
    ])
    rows_table(doc, "二十四、第十一章业务题二：银行存款余额调节表", ["项目", "金额"], [
        ("企业银行存款日记账余额", "80,000"),
        ("加：银行已收、企业未收", "50,000"),
        ("减：银行已付、企业未付", "5,000"),
        ("调节后余额", "125,000"),
        ("银行对账单余额", "75,000"),
        ("加：企业已收、银行未收", "52,000"),
        ("减：企业已付、银行未付", "2,000"),
        ("调节后余额", "125,000"),
    ])

    entries11 = [
        ("1", ["借：待处理财产损溢 15,000", "借：累计折旧 45,000", "贷：固定资产 60,000", "借：营业外支出 15,000", "贷：待处理财产损溢 15,000"]),
        ("2", ["借：固定资产 5,000", "贷：累计折旧 2,500", "贷：待处理财产损溢 2,500", "借：待处理财产损溢 2,500", "贷：营业外收入 2,500"]),
        ("3", ["借：库存现金 100", "贷：待处理财产损溢 100", "借：待处理财产损溢 100", "贷：营业外收入 100"]),
        ("4", ["盘亏金额 = (300 - 292) x 20 = 160", "借：待处理财产损溢 160", "贷：原材料-甲材料 160", "借：其他应收款 160", "贷：待处理财产损溢 160"]),
        ("5", ["盘盈金额 = (460 - 450) x 15 = 150", "借：原材料-乙材料 150", "贷：待处理财产损溢 150", "借：待处理财产损溢 150", "贷：管理费用 150"]),
        ("6", ["借：待处理财产损溢 25,000", "贷：原材料-丙材料 25,000", "借：管理费用 2,000", "借：其他应收款 18,000", "借：营业外支出 5,000", "贷：待处理财产损溢 25,000"]),
        ("7", ["盘亏金额 = (490 - 480) x 20 = 200", "借：待处理财产损溢 200", "贷：原材料-丁材料 200", "借：管理费用 200", "贷：待处理财产损溢 200"]),
        ("8", ["借：待处理财产损溢 5,000", "贷：库存商品-A产品 5,000", "借：库存现金 1,000", "借：其他应收款 1,500", "借：管理费用 2,500", "贷：待处理财产损溢 5,000"]),
        ("9", ["借：坏账准备 2,000", "贷：应收账款 2,000"]),
        ("10", ["应补提坏账准备 = 38,000 - 20,000 = 18,000", "借：资产减值损失 18,000", "贷：坏账准备 18,000"]),
    ]
    rows_table(doc, "二十五、第十一章业务题三：会计分录", ["题号", "答案"], entry_rows(entries11))
    doc.save(OUT / "汪政扬_第八次作业.docx")


if __name__ == "__main__":
    build_hw7()
    build_hw8()
    print(OUT / "汪政扬_第七次作业.docx")
    print(OUT / "汪政扬_第八次作业.docx")
