from __future__ import annotations

import re
import shutil
from collections import Counter, defaultdict
from pathlib import Path

import jieba
import jieba.posseg as pseg
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


ROOT = Path(r"C:\Users\28026\xwechat_files\wxid_6p586amrjhon22_4737\msg\file\2026-07\秘书")
OUT_ROOT = ROOT / "成品"
YEAR = "2025"
SUBJECT = "秘书学"
JOURNAL = "秘书之友"

STOPWORDS = {
    "的", "了", "和", "与", "及", "在", "对", "把", "被", "为", "以", "等", "中", "上", "下", "要",
    "是", "有", "也", "就", "都", "而", "或", "一个", "一种", "通过", "进行", "不断", "同时",
    "以及", "对于", "方面", "工作", "相关", "可以", "需要", "必须", "我们", "他们", "自己",
    "这个", "这些", "这种", "由于", "因此", "并且", "如果", "作为", "其中", "主要", "进一步",
    "进行", "实现", "推进", "提升", "加强", "不断", "通过", "具有", "形成", "推动", "做好",
    "问题", "方式", "情况", "过程", "内容", "作用", "要求", "发展", "建设", "研究",
}

DOMAIN_TERMS = [
    "秘书学", "秘书工作", "办公室", "三服务", "公文写作", "公文处理", "督查督办", "调查研究",
    "会议服务", "文稿写作", "信息工作", "档案管理", "党政秘书", "教学秘书", "科研秘书",
    "行政管理", "基层治理", "人工智能", "生成式人工智能", "DeepSeek", "AIGC", "智慧办公",
    "电子公文", "请示报告", "办文指数", "党政机关", "高校办公室", "高职院校", "应用文写作",
    "党政公文", "文电工作", "文秘写作", "秘书实践", "秘书素养", "文稿服务", "公务接待",
    "会议报到", "会议摄影", "视频会议", "讲话稿", "发言稿", "咨政报告", "申论归纳概括题",
    "基层减负", "第一议题", "三步工作法", "四轮驱动", "四定四化四融合", "清单制",
    "标准操作规范", "BOPPPS模型", "金课标准", "课程思政", "政产学研用", "研究性教学",
    "办公自动化", "智慧办公", "流动办公", "知识管理", "数字化转型", "AI赋能",
    "AI幻觉", "人机协作", "认知升维", "提示词", "大语言模型", "电子公文办理",
    "党委信息", "督查工作", "督办机制", "督办工作", "复核审查", "结果运用",
    "红黄绿预警", "一单两函", "督办通知单", "督办提示函", "督办警示函",
]


def decode_text(path: Path) -> str:
    data = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def normalize_fragment(text: str) -> str:
    text = text.replace("\u3000", " ")
    text = re.sub(r"[ \t]+", " ", text.strip())
    text = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", text)
    text = re.sub(r"(?<=[A-Za-z])\s+(?=[A-Za-z])", "", text)
    text = re.sub(r"(?<=\d)\s+(?=\d)", "", text)
    text = re.sub(r"\s+([，。；：！？、）】》])", r"\1", text)
    text = re.sub(r"([（【《])\s+", r"\1", text)
    text = text.replace("［］", "［J］").replace("J .", "J].")
    text = re.sub(r"(?<=[\u4e00-\u9fff])\s*[-－—]\s*(?=[\u4e00-\u9fff])", "", text)
    text = re.sub(r"[�□■●◆◇▲△▼▽]", "", text)
    return text.strip()


def is_residual_fragment(text: str) -> bool:
    compact = re.sub(r"\s+", "", text)
    if not compact:
        return True
    if re.fullmatch(r"\d+", compact):
        return True
    if re.fullmatch(r"[图表]\d+([-.—－]\d+)?", compact):
        return True
    if compact.startswith(("资料来源", "数据来源", "续表", "续图", "注：")):
        return True
    if re.search(r"秘书之友20\d{2}", compact):
        return True
    return False


def split_columns(line: str) -> list[str]:
    raw = line.rstrip()
    if not raw.strip():
        return []
    compact = re.sub(r"\s+", "", raw)
    if compact in {"督查工作", "秘书工作", "公文写作", "调查研究", "会务工作", "档案工作", "写作研究"}:
        return []
    if re.search(r"^\d+\s*秘书之友\s*2025·\d+$", raw):
        return []
    if re.search(r"^秘书之友\s*2025·\d+\s*\d+$", raw):
        return []

    # Three-column journal extraction commonly has column gaps ending around
    # 30 and 60 characters. Split on those gaps if present; otherwise split on
    # very long whitespace runs. This keeps short PDF intra-word spaces repairable.
    runs = list(re.finditer(r"[ \t]{4,}", raw))
    first = next((m for m in runs if 28 <= m.end() <= 38), None)
    if first:
        second = next((m for m in runs if m.start() > first.end() and 56 <= m.end() <= 70), None)
        if second:
            parts = [raw[: first.start()], raw[first.end() : second.start()], raw[second.end() :]]
        else:
            rest = raw[first.end() :]
            soft_runs = [m for m in re.finditer(r"[ \t]{4,}", rest) if m.start() >= 18]
            if soft_runs:
                soft = soft_runs[-1]
                parts = [raw[: first.start()], rest[: soft.start()], rest[soft.end() :]]
            else:
                parts = [raw[: first.start()], rest]
    else:
        parts = re.split(r"[ \t]{12,}", raw)
    cleaned_parts = []
    for part in parts:
        frag = normalize_fragment(part)
        if frag and not is_residual_fragment(frag):
            cleaned_parts.append(frag)
    return cleaned_parts


def repair_final_text(text: str) -> str:
    text = text.replace("\r", "")
    text = re.sub(r"[�□■●◆◇▲△▼▽]", "", text)
    text = re.sub(r"(?<=[\u4e00-\u9fff])\s*[-－—]\s*(?=[\u4e00-\u9fff])", "", text)
    text = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", text)
    text = re.sub(r"(?<=[A-Za-z])\s+(?=[A-Za-z])", "", text)
    text = re.sub(r"(?<=\d)\s+(?=\d)", "", text)
    text = re.sub(r"\s+([，。；：！？、）】》])", r"\1", text)
    text = re.sub(r"([（【《])\s+", r"\1", text)
    # Remove common references and chart/table debris that survive column splitting.
    text = re.sub(r"参考文献[:：].*", "", text, flags=re.S)
    text = re.sub(r"（作者单位[:：]?.*?）", "", text)
    text = re.sub(r"［\d+］", "", text)
    text = re.sub(r"\[[0-9]+\]", "", text)
    text = re.sub(r"(?m)^\s*(图|表)\s*\d+([-.—－]\d+)?[^\n]{0,80}$", "", text)
    text = re.sub(r"(?m)^\s*(资料来源|数据来源|续表|续图|注：).*$", "", text)
    text = re.sub(r"(?m)^\s*\d+\s*$", "", text)
    # Normalize section headings and paragraph breaks.
    text = re.sub(r"([。！？])", r"\1\n", text)
    text = re.sub(r"(?<!\n)([一二三四五六七八九十]、)", r"\n\1", text)
    text = re.sub(r"\n{2,}", "\n", text)
    return text.strip() + "\n"


def clean_text(raw: str) -> str:
    raw = raw.replace("\r", "")
    lines = raw.split("\n")
    pages: list[list[list[str]]] = []
    current_cols = [[], [], []]
    ref_skip = [False, False, False]
    started = False
    skip_until_footer = False
    skip_table_block = False

    for line in lines:
        stripped = line.strip()
        if not stripped:
            skip_table_block = False
            continue
        if "■■■■" in stripped:
            skip_until_footer = True
            continue
        if skip_until_footer:
            if re.search(r"\d+\s*秘书之友\s*2025·\d+", stripped):
                skip_until_footer = False
            continue
        if re.search(r"^表\s*\d+", stripped):
            skip_table_block = True
            continue
        if skip_table_block:
            continue
        if re.search(r"秘书之友\s*2025·\d+|2025·\d+\s*$", stripped) and re.search(r"\d", stripped):
            if any(current_cols):
                pages.append(current_cols)
                current_cols = [[], [], []]
                ref_skip = [False, False, False]
            continue
        if not started:
            # Drop mastheads and title clutter until plausible body starts.
            if re.search(r"[。！？]|\b(一|二|三|四|五)、", stripped) and len(re.sub(r"\s+", "", stripped)) > 18:
                started = True
            else:
                continue
        parts = split_columns(line)
        for idx, part in enumerate(parts[:3]):
            if not part:
                continue
            if ref_skip[idx]:
                continue
            if "参考文献" in part or "作者单位" in part or re.match(r"^［\d+］", part):
                ref_skip[idx] = True
                continue
            current_cols[idx].append(part)
    if any(current_cols):
        pages.append(current_cols)

    pieces = []
    for cols in pages:
        page_text = "".join("".join(col) for col in cols if col)
        if page_text:
            pieces.append(page_text)
    text = "".join(pieces)
    return repair_final_text(text)


def tokenize(text: str) -> list[tuple[str, str]]:
    for term in DOMAIN_TERMS:
        jieba.add_word(term, freq=200000, tag="n")
    tokens = []
    for word, flag in pseg.cut(text):
        word = word.strip()
        if len(word) < 2:
            continue
        if word in STOPWORDS:
            continue
        if re.fullmatch(r"[\W_]+", word):
            continue
        if re.fullmatch(r"\d+|[A-Za-z]+", word):
            continue
        tokens.append((word, flag))
    return tokens


def write_frequency_xlsx(path: Path, tokens: list[tuple[str, str]], source_name: str) -> None:
    counts = Counter(word for word, _ in tokens)
    pos_counts: dict[str, Counter] = defaultdict(Counter)
    for word, flag in tokens:
        pos_counts[word][flag] += 1

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "词频表"
    ws.append(["排名", "词语", "词性标注", "词频", "来源文件"])
    for rank, (word, count) in enumerate(counts.most_common(), 1):
        pos = pos_counts[word].most_common(1)[0][0]
        ws.append([rank, word, pos, count, source_name])

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")
        row[0].alignment = Alignment(horizontal="center")
        row[2].alignment = Alignment(horizontal="center")
        row[3].alignment = Alignment(horizontal="center")

    widths = [8, 28, 12, 10, 52]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main() -> None:
    files = sorted(p for p in ROOT.glob("*.txt") if p.is_file())
    OUT_ROOT.mkdir(parents=True, exist_ok=True)

    for index, src in enumerate(files, 1):
        stem = f"{YEAR}_{SUBJECT}_{JOURNAL}_{index:03d}"
        folder = OUT_ROOT / stem
        if folder.exists():
            shutil.rmtree(folder)
        folder.mkdir(parents=True)

        raw = decode_text(src)
        cleaned = clean_text(raw)
        txt_path = folder / f"{stem}.txt"
        xlsx_path = folder / f"{stem}_词频词性表.xlsx"
        txt_path.write_text(cleaned, encoding="utf-8-sig")
        write_frequency_xlsx(xlsx_path, tokenize(cleaned), src.name)

    print(f"processed={len(files)}")
    print(str(OUT_ROOT))


if __name__ == "__main__":
    main()
